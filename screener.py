# -*- coding: utf-8 -*-
"""
=======================================================================
 스윙 스크리너 (웹사이트용)  -  GitHub Actions에서 매일 자동 실행
=======================================================================
 Colab 통합본(swing_screening_colab.py)과 같은 계산을 한 뒤, 결과를
 웹페이지가 읽을 수 있는 파일로 저장합니다.

   docs/data/latest.json   : 웹페이지(docs/index.html)가 읽는 데이터
   docs/data/swing_latest.xlsx : 같은 결과의 엑셀 파일 (웹페이지에서 다운로드)

 직접 실행:  python screener.py
 기준을 바꾸려면 아래 '설정' 숫자를 고치세요. 웹페이지의 조건 입력칸
 기본값도 이 숫자를 따릅니다.
=======================================================================
"""

# ======================= 설정 (필요하면 바꾸세요) =======================
VPA_대상 = "전부"   # VPA로 볼 종목: "전부"(1개월+2주 리더 합침) / "1개월" / "2주"
VPA_차트수 = 5       # 차트를 그릴 VPA 점수 상위 종목 수
시총_하한_억 = 3000   # 시가총액이 이 금액(억원) 이상인 종목만 사용. 0으로 두면 필터 없음
VPA_최소_1개월 = 10   # VPA 결과에 넣을 종목의 1개월 등락률 하한(%). 이 값 이상만 남김
VPA_최소_2주 = 0      # VPA 결과에 넣을 종목의 최근 2주 등락률 하한(%). 0 = 마이너스 종목 제외

# --- 눌림목 조건 (VPA 결과 4번 시트에 AND로 함께 적용) ---
VPA_후보 = "전체"     # "전체": 시총 조건을 통과한 전 종목에서 찾기 / "리더": 업종 리더종목(2·3번 시트)에서만 찾기
업종RS_하한 = 80      # 종목이 속한 업종의 업종RS(1개월) 하한 → 강세 업종 소속만
종목RS_하한 = 80      # 종목RS(1개월 등락률의 전 종목 중 백분위, 1~99) 하한
종목RS_상한 = 99      # 종목RS 상한 (99 = 상한 없음). 과열 종목을 빼려면 90 등으로 낮추세요
이격_상한 = 1.05      # 현재가 ÷ 20일 이동평균선 ≤ 1.05 (20일선 대비 5% 이내)
이격_하한 = 0         # 0이면 하한 없음. 1.0으로 두면 20일선 위에 있는 종목만
거래량비_상한 = 0.5   # 최근 3거래일 평균 거래량 ÷ 직전 20거래일 평균 거래량 ≤ 0.5 (No Supply)
# =======================================================================

import io
import os
import re
import time
import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
import pandas as pd
import numpy as np
import FinanceDataReader as fdr

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
from matplotlib.patches import Rectangle

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


PRICE_CACHE = {}  # 종목코드 -> 시세 DataFrame (1단계에서 받아 2단계에서 재사용)
MARKET_CAP = {}   # 종목코드 -> 시가총액(억원)
CAP_NOTE = "시가총액 필터 없음"

# ---------------------------------------------------------------------
# 차트용 한글 폰트
# ---------------------------------------------------------------------
_candidates = ["Malgun Gothic", "AppleGothic", "Apple SD Gothic Neo",
               "NanumGothic", "Noto Sans CJK KR", "Noto Sans KR", "Noto Sans CJK JP"]


def _find_korean_font():
    installed = {f.name for f in fm.fontManager.ttflist}
    return next((n for n in _candidates if n in installed), None)


def _ensure_korean_font():
    """Colab에는 한글 글꼴이 없어서 차트 속 종목명이 □로 깨진다 → 나눔글꼴을 자동 설치"""
    font = _find_korean_font()
    if font:
        return font
    import glob
    import subprocess
    for cmd in (["apt-get", "-qq", "install", "-y", "fonts-nanum"],
                ["apt-get", "-qq", "update"],
                ["apt-get", "-qq", "install", "-y", "fonts-nanum"]):
        if glob.glob("/usr/share/fonts/truetype/nanum/NanumGothic*.ttf"):
            break
        try:
            subprocess.run(cmd, capture_output=True, timeout=180)
        except Exception:
            pass
    for path in glob.glob("/usr/share/fonts/truetype/nanum/*.ttf"):
        try:
            fm.fontManager.addfont(path)
        except Exception:
            pass
    font = _find_korean_font()
    if not font:
        print("  한글 글꼴을 설치하지 못했습니다. 차트 속 한글이 □로 보일 수 있습니다 (엑셀 표는 정상).")
    return font


KOREAN_FONT = _ensure_korean_font()
if KOREAN_FONT:
    plt.rcParams["font.family"] = KOREAN_FONT
plt.rcParams["axes.unicode_minus"] = False


def _t(korean_text, english_text):
    """한글 폰트가 없는 환경이면 영문 라벨로 자동 대체"""
    return korean_text if KOREAN_FONT else english_text



# ---------------------------------------------------------------------
# 1단계: 업종 분류·종목 지표·신호
# ---------------------------------------------------------------------
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Referer": "https://finance.naver.com/",
}
MIN_MEMBERS = 3  # 종목이 3개 미만인 업종은 평균이 한두 종목에 좌우되므로 순위에서 제외


# ---------------------------------------------------------------------
# 1-①. 네이버 증권 업종 분류 (2026년 9월 개편 이후의 JSON API)
#   네이버 금융이 'Npay 증권'으로 개편되면서 예전 업종 HTML 페이지에는
#   더 이상 업종 링크가 들어있지 않다. 대신 화면이 내부적으로 쓰는
#   JSON 주소에서 같은 79개 업종과 소속 종목을 받아온다.
# ---------------------------------------------------------------------
NAVER_API = "https://m.stock.naver.com/api"
NAVER_HEADERS = {
    "User-Agent": HEADERS["User-Agent"],
    "Referer": "https://m.stock.naver.com/",
    "Accept": "application/json, text/plain, */*",
}


def _pick(d, keys):
    """dict에서 후보 키 중 처음으로 값이 있는 것을 돌려준다 (응답 필드명이 조금 달라도 대응)"""
    for k in keys:
        v = d.get(k)
        if v not in (None, ""):
            return v
    return None


def _naver_get(url):
    res = requests.get(url, headers=NAVER_HEADERS, timeout=10)
    try:
        return res.json()
    except Exception:
        raise RuntimeError(f"JSON이 아닌 응답 (응답코드 {res.status_code}): {res.text[:150]!r}")


def _naver_industry_members(no):
    out, page, size = [], 1, 100
    while True:
        data = _naver_get(f"{NAVER_API}/stocks/industry/{no}?page={page}&pageSize={size}")
        stocks = data.get("stocks") or []
        for s in stocks:
            code = str(_pick(s, ["itemCode", "stockCode", "symbolCode", "code", "reutersCode"]) or "").strip()
            name = _pick(s, ["stockName", "itemName", "stockNameKor", "name"])
            if len(code) == 6 and name:
                out.append({"종목코드": code, "종목명": str(name), "주요제품": "",
                            "시총(억)": _cap_from_obj(s)})
        total = data.get("totalCount") or 0
        if not stocks or len(stocks) < size or page * size >= total:
            break
        page += 1
    return out


def try_naver():
    try:
        data = _naver_get(f"{NAVER_API}/stocks/industry?page=1&pageSize=100")
    except Exception as e:
        print(f"    네이버 업종 목록 조회 실패: {e}")
        return None

    groups = data.get("groups") or []
    sectors = []
    for g in groups:
        no = _pick(g, ["no", "groupNo", "industryNo", "industryCode", "code", "id"])
        name = _pick(g, ["name", "groupName", "industryName", "industryGroupKor"])
        if no is not None and name:
            sectors.append((str(name), str(no)))

    if len(sectors) < 5:
        print(f"    네이버 업종 목록 형식이 예상과 다름 (최상위 키: {list(data)[:10]}, "
              f"첫 항목: {str(groups[0])[:200] if groups else '없음'})")
        return None

    result, errors = {}, []
    with ThreadPoolExecutor(max_workers=6) as ex:
        fmap = {ex.submit(_naver_industry_members, no): name for name, no in sectors}
        for f in as_completed(fmap):
            try:
                result[fmap[f]] = f.result()
            except Exception as e:
                errors.append(str(e))

    total = sum(len(v) for v in result.values())
    if total < 100:
        sample = errors[0] if errors else "오류 없음"
        print(f"    네이버 업종별 종목을 충분히 읽지 못함 (총 {total}개, 예: {sample[:200]})")
        return None
    return result


# ---------------------------------------------------------------------
# 1-②. 한국거래소 KIND 상장법인목록
# ---------------------------------------------------------------------
def _group_by_sector(df, code_col, name_col, sector_col, product_col=None):
    df = df.dropna(subset=[code_col, sector_col]).copy()
    df[code_col] = df[code_col].astype(str).str.strip().str.replace(".0", "", regex=False).str.zfill(6)
    result = {}
    for _, row in df.iterrows():
        sector = str(row[sector_col]).strip()
        if not sector or sector.lower() == "nan":
            continue
        result.setdefault(sector, []).append({
            "종목코드": row[code_col],
            "종목명": str(row[name_col]).strip(),
            "주요제품": str(row[product_col]).strip() if product_col and pd.notna(row[product_col]) else "",
        })
    return result


def try_kind():
    url = "https://kind.krx.co.kr/corpgeneral/corpList.do?method=download&searchType=13"
    try:
        res = requests.get(url, headers={"User-Agent": HEADERS["User-Agent"]}, timeout=20)
        html = res.content.decode("euc-kr", errors="replace")
        df = pd.read_html(io.StringIO(html), header=0)[0]
    except Exception as e:
        print(f"    KIND 조회 실패: {e}")
        return None

    need = {"회사명", "종목코드", "업종"}
    if not need.issubset(df.columns):
        print(f"    KIND 표 형식이 예상과 다름 (컬럼: {list(df.columns)})")
        return None
    if "시장구분" in df.columns:  # 코넥스 제외
        df = df[df["시장구분"].astype(str).str.contains("유가|코스닥")]
    product = "주요제품" if "주요제품" in df.columns else None
    result = _group_by_sector(df, "종목코드", "회사명", "업종", product)
    return result if len(result) >= 5 else None


# ---------------------------------------------------------------------
# 1-③. FinanceDataReader KRX-DESC
# ---------------------------------------------------------------------
def try_fdr_desc():
    try:
        df = fdr.StockListing("KRX-DESC")
    except Exception as e:
        print(f"    FinanceDataReader KRX-DESC 조회 실패: {e}")
        return None
    code_col = "Code" if "Code" in df.columns else ("Symbol" if "Symbol" in df.columns else None)
    if code_col is None or "Sector" not in df.columns:
        print(f"    KRX-DESC 형식이 예상과 다름 (컬럼: {list(df.columns)})")
        return None
    if "Market" in df.columns:
        df = df[df["Market"].astype(str).str.upper().str.contains("KOSPI|KOSDAQ")]
    product = "Industry" if "Industry" in df.columns else None
    result = _group_by_sector(df, code_col, "Name", "Sector", product)
    return result if len(result) >= 5 else None


def get_sector_members():
    sources = [
        ("네이버 증권 업종 분류", try_naver),
        ("한국거래소 KIND 업종(통계청 산업분류)", try_kind),
        ("FinanceDataReader KRX-DESC 업종(통계청 산업분류)", try_fdr_desc),
    ]
    for label, fn in sources:
        print(f"  - {label} 시도 중...")
        result = fn()
        if result:
            print(f"    성공: 업종 {len(result)}개")
            return result, label
    raise RuntimeError(
        "업종 분류를 세 곳 모두에서 가져오지 못했습니다. 위에 출력된 '실패' 메시지를 "
        "그대로 복사해서 알려주세요."
    )


# ---------------------------------------------------------------------
# 2. 개별 종목 지표: 1개월/2주/1주 등락률 + 거래량배율
# ---------------------------------------------------------------------
def parse_eok(v):
    """'2,069조 5,826억', '3,123억원' 같은 글자를 억원 숫자로. 단위가 없는 숫자는 단위를 몰라서 None"""
    if v is None or isinstance(v, (int, float)):
        return None
    t = str(v).replace(",", "").replace(" ", "").replace("원", "")
    jo = re.search(r"([\d.]+)조", t)
    eok = re.search(r"([\d.]+)억", t)
    if not jo and not eok:
        return None
    return (float(jo.group(1)) * 10000 if jo else 0) + (float(eok.group(1)) if eok else 0)


def _cap_from_obj(obj):
    for k, v in obj.items():
        if any(t in k.lower() for t in ("marketvalue", "marketcap", "marketsum")):
            val = parse_eok(v)
            if val:
                return val
    return None


def _naver_cap(code):
    try:
        data = _naver_get(f"{NAVER_API}/stock/{code}/integration")
    except Exception:
        return None
    for it in data.get("totalInfos") or []:
        if it.get("code") == "marketValue" or "시총" in str(it.get("key", "")) or "시가총액" in str(it.get("key", "")):
            return parse_eok(it.get("value"))
    return None


def get_market_caps(codes, known):
    """시가총액(억원)을 ① 업종 목록 응답 ② FinanceDataReader 상장목록 ③ 네이버 종목별 조회 순으로 채운다"""
    caps = {c: v for c, v in known.items() if c in set(codes) and v}
    sources = ["네이버 증권 업종 목록"] if caps else []
    missing = [c for c in codes if c not in caps]

    if len(missing) > len(codes) * 0.1:
        try:
            lst = fdr.StockListing("KRX")
            ccol = "Code" if "Code" in lst.columns else "Symbol"
            got = 0
            for code, marcap in zip(lst[ccol].astype(str).str.zfill(6), lst["Marcap"]):
                if code in missing and pd.notna(marcap) and marcap > 0:
                    caps[code] = marcap / 1e8
                    got += 1
            if got:
                sources.append("FinanceDataReader 상장목록")
        except Exception as e:
            print(f"    FinanceDataReader 상장목록에서 시가총액 조회 실패: {e}")
        missing = [c for c in codes if c not in caps]

    if len(missing) > len(codes) * 0.1:
        print(f"    네이버에서 종목별 시가총액 조회 중... ({len(missing)}개, 1~2분 소요)")
        got = 0
        with ThreadPoolExecutor(max_workers=12) as ex:
            fmap = {ex.submit(_naver_cap, c): c for c in missing}
            for f in as_completed(fmap):
                v = f.result()
                if v:
                    caps[fmap[f]] = v
                    got += 1
        if got:
            sources.append("네이버 종목별 조회")

    print(f"    시가총액 확인: {len(caps)}/{len(codes)}개 종목")
    return caps, " + ".join(sources) if sources else "없음"


def _pct(a, b):
    if a is None or b is None or pd.isna(a) or pd.isna(b) or a == 0:
        return None
    return round((b - a) / a * 100, 2)


def get_stock_metrics(code, fetch_from, month_from, todate):
    try:
        df = fdr.DataReader(code, fetch_from, todate)
    except Exception:
        return None
    if df is None or len(df) < 6:
        return None
    PRICE_CACHE[code] = df  # VPA 패턴 스크리닝에서 재사용
    close, vol = df["Close"], df["Volume"]
    month = df[df.index >= pd.Timestamp(month_from)]
    out = {
        "1개월": _pct(month["Close"].iloc[0], month["Close"].iloc[-1]) if len(month) >= 2 else None,
        "2주": _pct(close.iloc[-11], close.iloc[-1]) if len(df) >= 11 else None,
        # 직전 2주: 1개월 시작일 ~ 최근 2주가 시작되기 직전 (한 달의 앞 절반)
        "직전2주": (_pct(month["Close"].iloc[0], close.iloc[-11])
                   if len(df) >= 11 and len(month) >= 2 and close.index[-11] > month.index[0] else None),
        "1주": _pct(close.iloc[-6], close.iloc[-1]),
        "거래량배율": None,
    }
    if len(df) >= 25:
        base = vol.iloc[-25:-5].mean()
        if base and base > 0:
            out["거래량배율"] = round(vol.iloc[-5:].mean() / base, 2)
    # 20일 이동평균선과 이격도 (현재가 ÷ 20일선)
    out["현재가"] = float(close.iloc[-1])
    out["20일선"] = round(float(close.iloc[-20:].mean()), 1) if len(df) >= 20 else None
    out["20일선 이격"] = round(out["현재가"] / out["20일선"], 3) if out["20일선"] else None
    # 최근 3거래일 평균 거래량 ÷ 그 직전 20거래일 평균 거래량
    out["거래량비(3일/20일)"] = None
    if len(df) >= 23:
        base20 = vol.iloc[-23:-3].mean()
        if base20 and base20 > 0:
            out["거래량비(3일/20일)"] = round(float(vol.iloc[-3:].mean() / base20), 2)
    return out if out["1개월"] is not None else None


# ---------------------------------------------------------------------
# 3. 업종 신호
# ---------------------------------------------------------------------
MIN_2W_GAIN = 2.0  # 순위만 높고 실제로는 거의 안 오른 업종을 거르기 위한 최소 2주 평균등락률(%)


def make_signal(rs1m, rs_prev, rs2w, rs1w, vol, ret2w):
    rising = ret2w is not None and not pd.isna(ret2w) and ret2w >= MIN_2W_GAIN
    if rising and rs2w >= 80 and rs_prev < 60:
        label = "신규 부상"
    elif rising and rs1m >= 80 and rs2w >= 80:
        label = "지속 강세"
    elif rs1m >= 80 and rs1w < 50:
        label = "강세 둔화"
    else:
        label = ""
    if vol is not None and not pd.isna(vol) and vol >= 1.5:
        label = f"{label} · 거래량↑" if label else "거래량↑"
    return label


# ---------------------------------------------------------------------
# 2단계: VPA 패턴 탐지
# ---------------------------------------------------------------------
def compute_indicators(df: "pd.DataFrame") -> "pd.DataFrame":
    df = df.copy()
    df["range_"] = df["High"] - df["Low"]
    df["body"] = (df["Close"] - df["Open"]).abs()
    df["atr14"] = df["range_"].rolling(14, min_periods=5).mean()
    df["ema5"] = df["Close"].ewm(span=5, adjust=False).mean()
    df["vol_ma5_prior"] = df["Volume"].shift(1).rolling(5, min_periods=3).mean()
    df["vol_ratio"] = df["Volume"] / df["vol_ma5_prior"]
    df["close_pos"] = (df["Close"] - df["Low"]) / df["range_"].replace(0, np.nan)
    df["lower_wick"] = df[["Open", "Close"]].min(axis=1) - df["Low"]
    df["lower_wick_ratio"] = df["lower_wick"] / df["range_"].replace(0, np.nan)
    df["swing_low_prior"] = df["Low"].shift(3).rolling(20, min_periods=10).min()
    return df


def _clip01(x):
    return max(0.0, min(1.0, float(x)))


def _trend_score(df, i, direction="down", lookback=5):
    if i - lookback < 0:
        return 0.0
    slope = df["ema5"].iloc[i - 1] - df["ema5"].iloc[i - lookback]
    ref = df["Close"].iloc[i - lookback] * 0.03
    if ref <= 0:
        return 0.0
    if direction == "down":
        return _clip01(-slope / ref)
    return _clip01(slope / ref)


def _score_no_supply_demand(df, i):
    row = df.iloc[i]
    if pd.isna(row["atr14"]) or pd.isna(row["vol_ratio"]) or row["atr14"] == 0:
        return None
    range_score = _clip01(1 - (row["range_"] / row["atr14"]) / 0.6)
    vol_score = _clip01((0.7 - row["vol_ratio"]) / 0.7)
    down_trend = _trend_score(df, i, "down")
    up_trend = _trend_score(df, i, "up")
    no_supply = (range_score + vol_score + down_trend) / 3
    no_demand = (range_score + vol_score + up_trend) / 3
    if no_supply >= no_demand:
        return ("No Supply", no_supply)
    return ("No Demand", no_demand)


def _score_stopping_volume(df, i):
    row = df.iloc[i]
    if pd.isna(row["vol_ratio"]) or pd.isna(row["close_pos"]):
        return None
    vol_score = _clip01((row["vol_ratio"] - 1) / 1.5)
    wick_score = _clip01((row["lower_wick_ratio"] - 0.3) / 0.4) if not pd.isna(row["lower_wick_ratio"]) else 0
    pos_score = _clip01((row["close_pos"] - 0.5) / 0.4)
    down_trend = _trend_score(df, i, "down")
    score = (vol_score + wick_score + pos_score + down_trend) / 4
    return ("Stopping Volume", score)


def _score_effort_result(df, i):
    row = df.iloc[i]
    if pd.isna(row["atr14"]) or pd.isna(row["vol_ratio"]) or row["atr14"] == 0:
        return None
    vol_score = _clip01((row["vol_ratio"] - 1) / 1.5)
    range_score = _clip01(1 - (row["range_"] / row["atr14"]) / 0.7)
    score = (vol_score + range_score) / 2
    return (_t("Effort-Result 불일치", "Effort-Result Mismatch"), score)


def _score_test_bar(df, i):
    row = df.iloc[i]
    if pd.isna(row["swing_low_prior"]) or pd.isna(row["vol_ratio"]) or row["swing_low_prior"] == 0:
        return None
    swing_low = row["swing_low_prior"]
    dist = (row["Low"] - swing_low) / swing_low
    if dist < -0.01:
        dist_score = _clip01(1 + dist / 0.02)
    else:
        dist_score = _clip01(1 - abs(dist) / 0.02)
    vol_score = _clip01((0.7 - row["vol_ratio"]) / 0.7)
    bounce_score = 1.0 if row["Close"] > row["Open"] else 0.4
    consolidation = _clip01(1 - _trend_score(df, i, "down"))
    score = (dist_score + vol_score + bounce_score) / 3 * (0.5 + 0.5 * consolidation)
    return ("Test Bar", score)


def best_pattern_in_window(df, window=5):
    """최근 window 거래일 중 4대 VPA 신호와 가장 유사한 캔들 하나를 찾는다."""
    df = compute_indicators(df)
    n = len(df)
    candidates = []
    for i in range(max(0, n - window), n):
        for fn in [_score_no_supply_demand, _score_stopping_volume,
                   _score_effort_result, _score_test_bar]:
            r = fn(df, i)
            if r is not None:
                name, score = r
                candidates.append((score, name, df.index[i], i))
    if not candidates:
        return None
    candidates.sort(key=lambda x: x[0], reverse=True)
    return candidates[0]


UP_COLOR, DOWN_COLOR = "#26A69A", "#EF5350"
VOL_COLOR, VOL_HILITE = "#B0BEC5", "#FF7043"
HILITE_EDGE = "#1F4E78"


def draw_pattern_chart(df, matched_date, pattern, score, ticker_label, out_path, lookback=30):
    df = df.tail(lookback).copy()
    df = df.reset_index().rename(columns={df.index.name or "index": "Date"})
    fig, (ax1, ax2) = plt.subplots(
        2, 1, figsize=(7.5, 4.2), dpi=150,
        gridspec_kw={"height_ratios": [3, 1], "hspace": 0.08}
    )
    for ax in (ax1, ax2):
        ax.set_xticks([])
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
    ax1.set_yticks([]); ax2.set_yticks([])

    matched_i = None
    for i, d in enumerate(df["Date"]):
        if hasattr(d, "date") and d.date() == matched_date.date():
            matched_i = i

    for i, row in df.iterrows():
        o, h, l, c, v = row["Open"], row["High"], row["Low"], row["Close"], row["Volume"]
        hl = (i == matched_i)
        color = UP_COLOR if c >= o else DOWN_COLOR
        ax1.plot([i, i], [l, h], color=color, linewidth=1.1, zorder=2)
        body_bottom, body_h = min(o, c), max(abs(c - o), (h - l) * 0.02)
        ax1.add_patch(Rectangle((i - 0.35, body_bottom), 0.7, body_h,
                                 facecolor=color,
                                 edgecolor=HILITE_EDGE if hl else color,
                                 linewidth=2.0 if hl else 0.7, zorder=3))
        vcolor = VOL_HILITE if hl else VOL_COLOR
        ax2.bar(i, v, width=0.7, color=vcolor,
                edgecolor=HILITE_EDGE if hl else vcolor,
                linewidth=1.4 if hl else 0, zorder=2)

    title = f"{ticker_label}  |  {pattern}  (점수 {score:.2f})"
    ax1.set_title(_t(title, title), fontsize=12.5, fontweight="bold", color=HILITE_EDGE, pad=10)
    ax2.set_ylabel(_t("거래량", "Volume"), fontsize=8.5, color="#666666")
    lo, hi = df["Low"].min(), df["High"].max()
    span = (hi - lo) or hi * 0.01
    ax1.set_ylim(lo - span * 0.03, hi + span * 0.22)  # 날짜 표시가 들어갈 여백
    if matched_i is not None:
        ax1.annotate(matched_date.strftime("%Y-%m-%d"),
                     xy=(matched_i, df["High"].iloc[matched_i]),
                     xytext=(matched_i, hi + span * 0.12),
                     fontsize=9, color=HILITE_EDGE, ha="center",
                     arrowprops=dict(arrowstyle="->", color=HILITE_EDGE, lw=1.2))
    fig.tight_layout()
    fig.savefig(out_path, facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------------
# 엑셀 시트 작성
# ---------------------------------------------------------------------
FILL = {
    "green": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "신규 부상": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    "지속 강세": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
    "강세 둔화": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}


def _write_table(ws, df, widths, row_fill_fn=None, signal_col=None):
    FONT_NAME = "Arial"
    cols = list(df.columns)
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if i - 1 < len(widths) else 11
    ws.row_dimensions[1].height = 30
    sig_idx = cols.index(signal_col) + 1 if signal_col in cols else None
    for r, rec in enumerate(df.to_dict("records"), start=2):
        fill = row_fill_fn(rec) if row_fill_fn else None
        for cidx, col in enumerate(cols, start=1):
            val = rec[col]
            if isinstance(val, float) and pd.isna(val):
                val = None
            cell = ws.cell(row=r, column=cidx, value=val)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.alignment = Alignment(horizontal="center")
            if fill:
                cell.fill = fill
        if sig_idx and rec.get(signal_col):
            key = rec[signal_col].split(" · ")[0]
            if key in FILL:
                ws.cell(row=r, column=sig_idx).fill = FILL[key]
    ws.freeze_panes = "C2"
    ws.sheet_view.showGridLines = False


def add_sector_sheets(wb, sector_df, leader_df, recent_df, source_label, month_from, todate):
    ws1 = wb.active
    ws1.title = "업종 순위"
    _write_table(ws1, sector_df, [6, 22, 8, 11, 9, 9, 11, 9, 11, 9, 11, 18],
                 row_fill_fn=lambda rec: FILL["green"] if rec["업종RS(1개월)"] >= 80 else None,
                 signal_col="신호")
    n = len(sector_df) + 3
    notes = [
        f"업종 분류 출처: {source_label}   |   1개월 기간: {month_from} ~ {todate}   |   2주=최근 10거래일, 1주=최근 5거래일, 직전2주=1개월 시작일~최근 2주 직전",
        f"대상 종목: {CAP_NOTE}",
        "평균등락률: 업종 소속 종목 등락률의 단순평균 (종목 3개 미만 업종 제외)   |   업종RS: 기간별 평균등락률의 업종 간 백분위(1~99)",
        "거래량배율: 최근 5거래일 평균 거래량 ÷ 그 전 20거래일 평균 거래량 (업종 내 중앙값). 1.0보다 크면 평소보다 거래가 늘어난 것",
        "신호  신규 부상: 2주RS≥80 & 직전2주RS<60  |  지속 강세: 1개월·2주RS 모두≥80  (둘 다 최근 2주 평균 +2% 이상일 때만)  |  강세 둔화: 1개월RS≥80 & 1주RS<50  |  거래량↑: 거래량배율≥1.5",
    ]
    for i, t in enumerate(notes):
        ws1.cell(row=n + i, column=2, value=t).font = Font(name="Arial", size=9, italic=True, color="808080")

    ws2 = wb.create_sheet("업종별 리더종목")
    _write_table(ws2, leader_df, [22, 10, 16, 10, 16, 12, 10, 10, 10, 10, 30], signal_col="업종 신호")

    ws3 = wb.create_sheet("최근 2주 강세 업종")
    _write_table(ws3, recent_df, [22, 10, 10, 10, 18, 10, 16, 12, 10, 10, 10, 10, 30], signal_col="업종 신호")



# ---------------------------------------------------------------------
# VPA 결과 시트
# ---------------------------------------------------------------------
def add_vpa_sheets(wb, vpa_df, chart_paths, top_n, note):
    ws = wb.create_sheet("VPA 스크리닝 결과")
    yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    _write_table(ws, vpa_df, [6, 10, 16, 20, 10, 18, 12, 8, 22, 8, 12, 12, 12, 11, 11, 11, 14],
                 row_fill_fn=lambda rec: yellow if rec["순위"] <= top_n and pd.notna(rec["점수"]) else None,
                 signal_col="업종 신호")
    n = len(vpa_df) + 3
    ws.cell(row=n, column=2,
            value=note + ("  →  조건을 만족하는 종목 없음" if len(vpa_df) == 0 else "")
            ).font = Font(name="Arial", size=9, italic=True, color="808080")
    n += 1
    ws.cell(row=n, column=2, value="점수: 최근 5거래일 캔들이 4가지 VPA 패턴 중 하나와 얼마나 비슷한지(0~1). "
            "1차 필터용 참고 지표이며 최종 진입은 차트로 직접 확인하세요.").font = Font(name="Arial", size=9, italic=True, color="808080")

    ws2 = wb.create_sheet("VPA 상위 차트")
    ws2.sheet_view.showGridLines = False
    ws2.column_dimensions["A"].width = 4
    row_cursor = 1
    for path in chart_paths:
        img = XLImage(path)
        img.width = img.width * 0.58
        img.height = img.height * 0.58
        ws2.add_image(img, f"B{row_cursor}")
        row_cursor += 24

# ---------------------------------------------------------------------
# 메인
# ---------------------------------------------------------------------
def _rs(series):
    return (series.rank(pct=True) * 98 + 1).round().astype("Int64")


def run_sector_stage(today):
    todate = today.strftime("%Y-%m-%d")
    month_from = (today - datetime.timedelta(days=30)).strftime("%Y-%m-%d")
    fetch_from = (today - datetime.timedelta(days=90)).strftime("%Y-%m-%d")  # VPA 지표 계산용으로 넉넉히
    print(f"1개월 기간: {month_from} ~ {todate} (2주·1주는 최근 10·5거래일)")

    print("\n[1/4] 업종 분류 가져오는 중...")
    sector_members, source_label = get_sector_members()
    all_codes = {}
    for members in sector_members.values():
        for m in members:
            all_codes[m["종목코드"]] = m["종목명"]
    print(f"  전체 종목(중복 제거): {len(all_codes)}개")

    global CAP_NOTE
    known = {m["종목코드"]: m.get("시총(억)") for ms in sector_members.values() for m in ms}
    if 시총_하한_억 and 시총_하한_억 > 0:
        print(f"\n  시가총액 확인 중... (기준: {시총_하한_억:,}억원 이상)")
        caps, cap_src = get_market_caps(list(all_codes), known)
        MARKET_CAP.update(caps)
        coverage = len(caps) / max(len(all_codes), 1)
        if coverage < 0.5:
            CAP_NOTE = f"시가총액을 {coverage:.0%} 종목만 확인해 필터를 적용하지 않음"
            print(f"  ! {CAP_NOTE}. 모든 종목으로 계속 진행합니다.")
        else:
            keep = {c for c, v in caps.items() if v >= 시총_하한_억}
            sector_members = {name: [m for m in ms if m["종목코드"] in keep]
                              for name, ms in sector_members.items()}
            unknown = len(all_codes) - len(caps)
            all_codes = {c: n for c, n in all_codes.items() if c in keep}
            CAP_NOTE = f"시가총액 {시총_하한_억:,}억원 이상 종목만 사용 (출처: {cap_src}, 조회 시점 기준)"
            print(f"  {시총_하한_억:,}억원 이상: {len(all_codes)}개 종목만 사용"
                  + (f" (시가총액 미확인 {unknown}개 제외)" if unknown else ""))
    else:
        MARKET_CAP.update({c: v for c, v in known.items() if v})

    print("\n[2/4] 종목별 시세 조회 및 등락률·거래량 계산 중... (수 분 소요될 수 있습니다)")
    t0 = time.time()
    metrics = {}
    with ThreadPoolExecutor(max_workers=12) as ex:
        fmap = {ex.submit(get_stock_metrics, c, fetch_from, month_from, todate): c for c in all_codes}
        done = 0
        for f in as_completed(fmap):
            metrics[fmap[f]] = f.result()
            done += 1
            if done % 200 == 0 or done == len(all_codes):
                print(f"    진행: {done}/{len(all_codes)} ({time.time()-t0:.0f}초 경과)")

    print("\n[3/4] 업종별 집계 및 순위 산출 중...")
    sector_rows, sector_detail = [], {}
    for name, members in sector_members.items():
        detail = []
        for m in members:
            mt = metrics.get(m["종목코드"])
            if mt:
                detail.append({"종목코드": m["종목코드"], "종목명": m["종목명"],
                               "시가총액(억)": MARKET_CAP.get(m["종목코드"]),
                               "주요제품": m.get("주요제품", ""), **mt})
        if len(detail) < MIN_MEMBERS:
            continue
        d = pd.DataFrame(detail)
        sector_detail[name] = d
        sector_rows.append({
            "업종명": name,
            "구성종목수": len(d),
            "1개월 평균등락률(%)": round(d["1개월"].mean(), 2),
            "직전2주 평균등락률(%)": round(d["직전2주"].mean(), 2) if d["직전2주"].notna().any() else None,
            "2주 평균등락률(%)": round(d["2주"].mean(), 2),
            "1주 평균등락률(%)": round(d["1주"].mean(), 2),
            "거래량배율": round(d["거래량배율"].median(), 2) if d["거래량배율"].notna().any() else None,
        })
    if not sector_rows:
        raise RuntimeError("등락률을 계산할 수 있는 업종이 없습니다. 인터넷 연결을 확인해주세요.")

    sdf = pd.DataFrame(sector_rows)
    sdf["업종RS(1개월)"] = _rs(sdf["1개월 평균등락률(%)"])
    sdf["업종RS(직전2주)"] = _rs(sdf["직전2주 평균등락률(%)"])
    sdf["업종RS(2주)"] = _rs(sdf["2주 평균등락률(%)"])
    sdf["업종RS(1주)"] = _rs(sdf["1주 평균등락률(%)"])
    sdf["신호"] = [make_signal(r["업종RS(1개월)"],
                               r["업종RS(직전2주)"] if pd.notna(r["업종RS(직전2주)"]) else 99,
                               r["업종RS(2주)"] if pd.notna(r["업종RS(2주)"]) else 0,
                               r["업종RS(1주)"] if pd.notna(r["업종RS(1주)"]) else 0,
                               r["거래량배율"], r["2주 평균등락률(%)"])
                  for _, r in sdf.iterrows()]
    sdf = sdf.sort_values("업종RS(1개월)", ascending=False).reset_index(drop=True)
    sdf.insert(0, "순위", range(1, len(sdf) + 1))
    sector_df = sdf[["순위", "업종명", "구성종목수",
                     "1개월 평균등락률(%)", "업종RS(1개월)", "업종RS(직전2주)",
                     "2주 평균등락률(%)", "업종RS(2주)",
                     "1주 평균등락률(%)", "업종RS(1주)",
                     "거래량배율", "신호"]]
    info = sdf.set_index("업종명")

    def leaders(sector_names, sort_col, rs_cols):
        rows = []
        for name in sector_names:
            top5 = sector_detail[name].sort_values(sort_col, ascending=False, na_position="last").head(5)
            for _, s in top5.iterrows():
                row = {"업종명": name}
                for c in rs_cols:
                    row[c] = info.loc[name, c]
                row.update({"업종 신호": info.loc[name, "신호"],
                            "종목코드": s["종목코드"], "종목명": s["종목명"],
                            "시가총액(억)": round(s["시가총액(억)"]) if pd.notna(s["시가총액(억)"]) else None})
                for c in [sort_col] + [x for x in ["1개월", "2주", "1주"] if x != sort_col]:
                    row[f"{c} 등락률(%)"] = s[c]
                row["거래량배율"] = s["거래량배율"]
                row["주요제품"] = s["주요제품"]
                rows.append(row)
        return pd.DataFrame(rows)

    top1m = sdf[sdf["업종RS(1개월)"] >= 80]
    if len(top1m) < 10:
        top1m = sdf.head(10)
    leader_df = leaders(top1m["업종명"], "1개월", ["업종RS(1개월)"])

    top2w = sdf[sdf["업종RS(2주)"] >= 80].sort_values("업종RS(2주)", ascending=False)
    if len(top2w) < 10:
        top2w = sdf.sort_values("업종RS(2주)", ascending=False).head(10)
    recent_df = leaders(top2w["업종명"], "2주", ["업종RS(2주)", "업종RS(직전2주)", "업종RS(1개월)"])

    print("  1개월 기준 상위 5개 업종:")
    for _, r in sector_df.head(5).iterrows():
        print(f"    {r['순위']}. {r['업종명']}  (1개월RS {r['업종RS(1개월)']}, 2주RS {r['업종RS(2주)']}) {r['신호']}")
    new = sector_df[sector_df["신호"].str.startswith("신규 부상")]
    print(f"  최근 2주 새로 부상한 업종: {len(new)}개")
    for _, r in new.sort_values("업종RS(2주)", ascending=False).head(10).iterrows():
        print(f"    - {r['업종명']}  (최근2주RS {r['업종RS(2주)']}, 직전2주RS {r['업종RS(직전2주)']}) {r['신호']}")

    # 전 종목 지표표: 종목 RS(1개월 등락률의 전 종목 중 백분위)와 20일선·거래량비 포함
    stock_rows, seen = [], set()
    for name, members in sector_members.items():
        for m in members:
            code = m["종목코드"]
            mt = metrics.get(code)
            if not mt or code in seen:
                continue
            seen.add(code)
            stock_rows.append({"종목코드": code, "종목명": m["종목명"], "업종명": name,
                               "업종RS(1개월)": info.loc[name, "업종RS(1개월)"] if name in info.index else None,
                               "업종 신호": info.loc[name, "신호"] if name in info.index else "",
                               "시가총액(억)": round(MARKET_CAP[code]) if MARKET_CAP.get(code) else None,
                               **{k: mt[k] for k in ["1개월", "2주", "1주", "현재가", "20일선",
                                                     "20일선 이격", "거래량비(3일/20일)"]}})
    all_df = pd.DataFrame(stock_rows).rename(columns={"1개월": "1개월 등락률(%)", "2주": "2주 등락률(%)",
                                                      "1주": "1주 등락률(%)"})
    all_df.insert(3, "종목RS", _rs(all_df["1개월 등락률(%)"]))
    all_df = all_df.sort_values("종목RS", ascending=False).reset_index(drop=True)

    return sector_df, leader_df, recent_df, all_df, source_label, month_from, todate


def run_vpa_stage(leader_df, recent_df, all_df):
    # 후보 모집단
    if VPA_후보 == "리더":
        if VPA_대상 == "1개월":
            codes = leader_df["종목코드"]
        elif VPA_대상 == "2주":
            codes = recent_df["종목코드"]
        else:
            codes = pd.concat([leader_df["종목코드"], recent_df["종목코드"]])
        cand = all_df[all_df["종목코드"].isin(set(codes))]
        pool = f"업종 리더종목({VPA_대상})"
    else:
        cand = all_df
        pool = "전 종목" + (f"(시총 {시총_하한_억:,}억원 이상)" if 시총_하한_억 else "")
    cand = cand.reset_index(drop=True)

    num = lambda col: pd.to_numeric(cand[col], errors="coerce")
    rs = num("종목RS")
    conds = [
        (f"1개월 등락률 ≥ {VPA_최소_1개월:g}%", num("1개월 등락률(%)") >= VPA_최소_1개월),
        (f"최근 2주 등락률 ≥ {VPA_최소_2주:g}%", num("2주 등락률(%)") >= VPA_최소_2주),
        (f"업종RS(1개월) ≥ {업종RS_하한}", num("업종RS(1개월)") >= 업종RS_하한),
        ((f"종목RS ≥ {종목RS_하한}" if 종목RS_상한 >= 99 else f"종목RS {종목RS_하한}~{종목RS_상한}"),
         (rs >= 종목RS_하한) & (rs <= 종목RS_상한)),
        (f"현재가/20일선 ≤ {이격_상한:g}" + (f" & ≥ {이격_하한:g}" if 이격_하한 else ""),
         (num("20일선 이격") <= 이격_상한) & (num("20일선 이격") >= (이격_하한 or 0))),
        (f"3일/20일 거래량비 ≤ {거래량비_상한:g}", num("거래량비(3일/20일)") <= 거래량비_상한),
    ]
    print(f"\n[4/4] 조건 필터 + VPA 패턴 스크리닝 (후보: {pool} {len(cand)}개)")
    mask = pd.Series(True, index=cand.index)
    for label, c in conds:
        mask &= c.fillna(False)
        print(f"    + {label:<28} → {int(mask.sum())}개")
    cand = cand[mask].reset_index(drop=True)
    if cand.empty:
        print("  조건을 모두 만족하는 종목이 없습니다. 코드 맨 위 설정값을 완화해보세요 "
              "(위 단계별 숫자를 보면 어느 조건에서 많이 걸러지는지 알 수 있습니다).")

    rows = []
    for _, c in cand.iterrows():
        code = c["종목코드"]
        base = {k: c[k] for k in ["종목코드", "종목명", "업종명", "업종RS(1개월)", "업종 신호", "시가총액(억)", "종목RS"]}
        base.update({"패턴": None, "점수": None, "매칭일": None})
        base.update({k: c[k] for k in ["1개월 등락률(%)", "2주 등락률(%)", "현재가", "20일선",
                                       "20일선 이격", "거래량비(3일/20일)"]})
        df = PRICE_CACHE.get(code)
        result = best_pattern_in_window(df) if df is not None and len(df) >= 10 else None
        if result is not None:
            score, pattern, matched_date, _ = result
            base.update({"패턴": pattern, "점수": round(float(score), 3),
                         "매칭일": matched_date.strftime("%Y-%m-%d")})
        rows.append(base)

    cols = ["종목코드", "종목명", "업종명", "업종RS(1개월)", "업종 신호", "시가총액(억)", "종목RS", "패턴", "점수", "매칭일",
            "1개월 등락률(%)", "2주 등락률(%)", "현재가", "20일선", "20일선 이격", "거래량비(3일/20일)"]
    vpa_df = pd.DataFrame(rows, columns=cols).sort_values("점수", ascending=False, na_position="last").reset_index(drop=True)
    vpa_df.insert(0, "순위", range(1, len(vpa_df) + 1))

    chart_paths = []
    os.makedirs("_charts_tmp", exist_ok=True)
    for _, r in vpa_df[vpa_df["점수"].notna()].head(VPA_차트수).iterrows():
        out_path = os.path.join("_charts_tmp", f"{r['종목코드']}.png")
        draw_pattern_chart(PRICE_CACHE[r["종목코드"]],
                           datetime.datetime.strptime(r["매칭일"], "%Y-%m-%d"),
                           r["패턴"], r["점수"], f"{r['종목코드']} {r['종목명']} [{r['업종명']}]", out_path)
        chart_paths.append(out_path)
        print(f"  {r['순위']}. {r['종목명']} [{r['업종명']}] 업종RS {r['업종RS(1개월)']}, 종목RS {r['종목RS']}, 20일선 이격 {r['20일선 이격']}, "
              f"거래량비 {r['거래량비(3일/20일)']} | {r['패턴']} (점수 {r['점수']:.2f})")
    cond_text = " & ".join(label for label, _ in conds)
    return vpa_df, chart_paths, f"후보: {pool}  |  조건: {cond_text}"



# ---------------------------------------------------------------------
# 웹페이지용 데이터 내보내기
# ---------------------------------------------------------------------
import json


def _clean(v):
    """JSON에 넣을 수 있게 NaN/넘파이 숫자를 정리"""
    if v is None:
        return None
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return None if pd.isna(v) else round(float(v), 4)
    if v is pd.NA:
        return None
    return v


def _records(df):
    return [{k: _clean(v) for k, v in rec.items()} for rec in df.to_dict("records")]


def score_all_vpa(all_df):
    """사이트에서 조건을 바꿔도 바로 보이도록, 전 종목의 VPA 패턴 점수를 미리 계산"""
    pats, scores, dates = [], [], []
    for code in all_df["종목코드"]:
        df = PRICE_CACHE.get(code)
        r = best_pattern_in_window(df) if df is not None and len(df) >= 10 else None
        if r is None:
            pats.append(None); scores.append(None); dates.append(None)
        else:
            s, p, d, _ = r
            pats.append(p); scores.append(round(float(s), 3)); dates.append(d.strftime("%Y-%m-%d"))
    out = all_df.copy()
    out["패턴"], out["점수"], out["매칭일"] = pats, scores, dates
    return out


def candles_for(code, n=40):
    df = PRICE_CACHE.get(code)
    if df is None or df.empty:
        return None
    ma20 = df["Close"].rolling(20).mean()
    t = df.tail(n)
    m = ma20.tail(n)
    return {
        "d": [x.strftime("%Y-%m-%d") for x in t.index],
        "o": [int(round(x)) for x in t["Open"]],
        "h": [int(round(x)) for x in t["High"]],
        "l": [int(round(x)) for x in t["Low"]],
        "c": [int(round(x)) for x in t["Close"]],
        "v": [int(x) for x in t["Volume"]],
        "ma": [None if pd.isna(x) else int(round(x)) for x in m],
    }


def main(out_dir="docs/data"):
    os.makedirs(out_dir, exist_ok=True)
    kst = datetime.timezone(datetime.timedelta(hours=9))
    now = datetime.datetime.now(kst)
    today = now.date()
    print("=" * 60)
    print(f"스윙 스크리너 (웹사이트용)  {now:%Y-%m-%d %H:%M} KST")
    print("=" * 60)

    sector_df, leader_df, recent_df, all_df, source_label, month_from, todate = run_sector_stage(today)
    vpa_df, chart_paths, vpa_note = run_vpa_stage(leader_df, recent_df, all_df)

    # 엑셀 (Colab 통합본과 같은 시트 구성)
    wb = openpyxl.Workbook()
    add_sector_sheets(wb, sector_df, leader_df, recent_df, source_label, month_from, todate)
    wb["업종별 리더종목"].title = "업종별 리더종목(1개월)"
    add_vpa_sheets(wb, vpa_df, chart_paths, VPA_차트수, vpa_note)
    ws6 = wb.create_sheet("전종목 지표")
    _write_table(ws6, all_df, [10, 16, 20, 8, 10, 18, 12, 11, 11, 11, 11, 11, 11, 14])
    xlsx_name = "swing_latest.xlsx"
    wb.save(os.path.join(out_dir, xlsx_name))

    # 웹페이지용 JSON
    print("\n웹페이지용 데이터 저장 중...")
    stocks = score_all_vpa(all_df)
    last_days = [df.index[-1] for df in PRICE_CACHE.values() if df is not None and len(df)]
    base_day = max(last_days).strftime("%Y-%m-%d") if last_days else todate  # 마지막 거래일
    data = {
        "meta": {
            "기준일": base_day, "생성시각": now.strftime("%Y-%m-%d %H:%M"), "1개월시작": month_from,
            "업종출처": source_label, "대상종목": CAP_NOTE, "엑셀": f"data/{xlsx_name}",
            "기본조건": {
                "VPA_후보": VPA_후보, "최소_1개월": VPA_최소_1개월, "최소_2주": VPA_최소_2주,
                "업종RS_하한": 업종RS_하한, "종목RS_하한": 종목RS_하한, "종목RS_상한": 종목RS_상한,
                "이격_상한": 이격_상한, "이격_하한": 이격_하한, "거래량비_상한": 거래량비_상한,
            },
        },
        "sectors": _records(sector_df),
        "leaders1m": _records(leader_df),
        "leaders2w": _records(recent_df),
        "stocks": _records(stocks),
        "candles": {code: candles_for(code) for code in stocks["종목코드"]},
    }
    path = os.path.join(out_dir, "latest.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    print(f"  {path} ({os.path.getsize(path)/1024:,.0f} KB), 종목 {len(stocks)}개")
    print("=" * 60)
    print("완료!")


if __name__ == "__main__":
    main()
